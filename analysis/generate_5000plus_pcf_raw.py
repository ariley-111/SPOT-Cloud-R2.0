"""Raw 5000+ Print+Cut+Finishing export for review (Q224297 die-cut hand-formed)."""
import os, shutil
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from collections import defaultdict

ANALYSIS_DIR = "/app/analysis"
STATIC_DIR = "/app/backend/static_assets"
OUT_FILE = "SPOT_5000plus_PCF_RawJoined.xlsx"

wb_src = openpyxl.load_workbook(os.path.join(STATIC_DIR, "SPOT_JobType_Variance_v3.xlsx"), data_only=False)
ws = wb_src["All Matched"]
h = [c.value for c in ws[1]]
matched = [dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True)]

wb_in = openpyxl.load_workbook(os.path.join(ANALYSIS_DIR, "input.xlsx"), data_only=True)
sp = wb_in["SPOT Data"]
sph = [c.value for c in sp[1]]
sp_rows = [dict(zip(sph, r)) for r in sp.iter_rows(min_row=2, values_only=True)]
sp_lookup = defaultdict(list)
for r in sp_rows:
    q = r.get("iQuote Estimate #")
    if q:
        sp_lookup[(q, r.get("Qty"))].append(r)

wb_full = openpyxl.load_workbook(os.path.join(ANALYSIS_DIR, "spot_full.xlsx"), data_only=True)
wf = wb_full["Testing Comparison"]
wfh = [c.value for c in wf[1]]
sf_lookup = defaultdict(list)
for r in wf.iter_rows(min_row=2, values_only=True):
    d = dict(zip(wfh, r))
    try: q = int(d.get("iQuote #") or 0) or None
    except (ValueError, TypeError): q = None
    if q:
        sf_lookup[(q, d.get("One-Up W"), d.get("One-Up L"))].append(d)

# Filter to 5000+ P+C+F
in_band = [r for r in matched if r["Qty Band"] == "5000+" and r["Job Type"] == "Print + Cut + Finishing"]
for r in in_band:
    candidates = sp_lookup.get((r["iQuote #"], r["Qty"]), [])
    best = min(candidates, key=lambda c: abs((c.get("Component Total") or 0) - (r["SPOT $"] or 0)), default=None)
    if best:
        r["W"] = best.get("Width"); r["L"] = best.get("Length")
        r["Board"] = best.get("Board"); r["Sheet"] = best.get("Sheet")
        r["PrintMode"] = best.get("Print Mode"); r["Route"] = best.get("Route")
    sf_rows = sf_lookup.get((r["iQuote #"], r.get("W"), r.get("L")), [])
    r["Print $"] = max((s.get("Print Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Cut $"] = max((s.get("Cutting Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Strip $"] = max((s.get("Strip Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Glue $"] = max((s.get("Glue Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Tape $"] = max((s.get("Tape Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Board $"] = max((s.get("Board Cost") or 0) for s in sf_rows) if sf_rows else 0
    r["Total Cost"] = max((s.get("Total Cost") or 0) for s in sf_rows) if sf_rows else 0

in_band.sort(key=lambda r: r["Qty"])

wb = openpyxl.Workbook()
ws_out = wb.active
ws_out.title = "5000+ PCF Raw"

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")

cols = [
    "iQuote #", "SPOT Estimate #", "Project", "Customer",
    "Job Type", "Component (SPOT)", "iQuote Product",
    "Qty", "W", "L", "Board", "Sheet", "Print Mode", "Route",
    "SPOT $", "iQuote $", "Δ $", "Δ %",
    "Print Cost", "Cut Cost", "Strip Cost", "Glue Cost", "Tape Cost", "Board Cost", "Total Cost",
]
for c, label in enumerate(cols, start=1):
    cell = ws_out.cell(row=1, column=c, value=label)
    cell.fill = HDR_FILL; cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

for i, r in enumerate(in_band, start=2):
    for c, key in enumerate(cols, start=1):
        if key == "Component (SPOT)": v = r.get("Component")
        elif key == "Print Cost": v = r.get("Print $")
        elif key == "Cut Cost": v = r.get("Cut $")
        elif key == "Strip Cost": v = r.get("Strip $")
        elif key == "Glue Cost": v = r.get("Glue $")
        elif key == "Tape Cost": v = r.get("Tape $")
        elif key == "Board Cost": v = r.get("Board $")
        else: v = r.get(key)
        cell = ws_out.cell(row=i, column=c, value=v)
        if key in ("SPOT $","iQuote $","Δ $","Print Cost","Cut Cost","Strip Cost","Glue Cost","Tape Cost","Board Cost","Total Cost") and isinstance(v,(int,float)):
            cell.number_format = '"$"#,##0.00'
        elif key == "Δ %" and isinstance(v,(int,float)):
            cell.number_format = "0.00%"

widths = [10, 20, 30, 26, 26, 30, 26, 8, 10, 10, 14, 18, 12, 10, 12, 12, 12, 10, 12, 12, 12, 12, 12, 12, 12]
for i, w in enumerate(widths, start=1):
    ws_out.column_dimensions[get_column_letter(i)].width = w
ws_out.freeze_panes = "A2"

local_path = os.path.join(ANALYSIS_DIR, OUT_FILE)
wb.save(local_path)
os.makedirs(STATIC_DIR, exist_ok=True)
shutil.copy(local_path, os.path.join(STATIC_DIR, OUT_FILE))
print(f"Wrote: {local_path}")
print(f"Rows exported: {len(in_band)}")
