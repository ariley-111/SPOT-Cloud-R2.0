"""Job-Type Variance analysis — CORRECTED join logic.

Replaces the ordinal-fallback pairing (which produced ~14% bad matches) with
a fuzzy name+qty matcher. Rows that cannot be matched with reasonable
confidence are marked unmatched and excluded from stats (not forced onto
arbitrary partners).

Match precedence:
  1. Exact (est, option, qty)
  2. (est, qty) when qty uniquely identifies one iQuote row for that estimate
  3. (est, qty ± 20%) + product-name similarity ≥ 0.55
  4. (est) + product-name similarity ≥ 0.65
  5. Unmatched (excluded)

Usage: python3 generate_jobtype_v2.py <full_spot.xlsx> <round2_input.xlsx> <output.xlsx>
"""
import sys
from difflib import SequenceMatcher

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SPOT_FULL = sys.argv[1] if len(sys.argv) > 1 else "spot_full.xlsx"
R2_INPUT  = sys.argv[2] if len(sys.argv) > 2 else "input.xlsx"
OUT       = sys.argv[3] if len(sys.argv) > 3 else "SPOT_JobType_Variance_v2.xlsx"


spf = pd.read_excel(SPOT_FULL, sheet_name="Testing Comparison")
iq  = pd.read_excel(R2_INPUT,  sheet_name="iQuote Data")
r2  = pd.read_excel(R2_INPUT,  sheet_name="SPOT Data")

r2_ids = set(r2["iQuote Estimate #"].dropna().astype(int).unique())
spf["iq_id"] = pd.to_numeric(spf["iQuote #"], errors="coerce").astype("Int64")
spf = spf[~spf["SPOT Estimate #"].astype(str).str.contains("-COPY", na=False)]
spf = spf[spf["iq_id"].isin(r2_ids)].reset_index(drop=True)


def norm_opt(x):
    if pd.isna(x): return ""
    s = str(x).strip().lower()
    if " - " in s: s = s.split(" - ")[0].strip()
    return s


def norm_name(x):
    if pd.isna(x): return ""
    s = str(x).strip().lower()
    # strip common qualifiers so "IBM (Display Body)" vs "IBM Display Body" collide cleanly
    for junk in ["—", "-", "(", ")", ":", ",", "/", "'", "\""]:
        s = s.replace(junk, " ")
    return " ".join(s.split())


def name_sim(a, b):
    return SequenceMatcher(None, norm_name(a), norm_name(b)).ratio()


spf["_opt"] = spf["iQuote Option"].apply(norm_opt)
iq["_opt"]  = iq["OptDesc"].apply(norm_opt)


# ─── Smarter matcher ─────────────────────────────────────────────
def match_iq(sp_row, iq_pool_available):
    est = sp_row["iq_id"]
    opt = sp_row["_opt"]
    qty = sp_row["Qty"]
    comp = sp_row["Component"]

    pool = iq_pool_available[iq_pool_available["IdQuote"] == est]
    if len(pool) == 0:
        return None, "no-iq-rows"

    # 1. exact (est, opt, qty)
    p = pool[(pool["_opt"] == opt) & (pool["Quantity"] == qty)]
    if len(p) >= 1:
        return p.index[0], "exact"

    # 2. (est, qty) when qty is unique for this estimate in iQuote
    p = pool[pool["Quantity"] == qty]
    if len(p) == 1:
        return p.index[0], "qty-unique"

    # If qty ties, use name similarity to disambiguate
    if len(p) >= 2 and pd.notna(comp):
        best = None; best_s = 0
        for idx, r in p.iterrows():
            s = name_sim(comp, r["Product"])
            if s > best_s:
                best_s, best = s, idx
        if best_s >= 0.55:
            return best, f"qty-tie+name({best_s:.2f})"

    # 3. qty within 20% + name similarity ≥ 0.55
    if pd.notna(qty) and pd.notna(comp):
        qmin, qmax = qty * 0.80, qty * 1.20
        p = pool[(pool["Quantity"] >= qmin) & (pool["Quantity"] <= qmax)]
        best = None; best_s = 0
        for idx, r in p.iterrows():
            s = name_sim(comp, r["Product"])
            if s > best_s:
                best_s, best = s, idx
        if best_s >= 0.55 and best is not None:
            return best, f"qty±20+name({best_s:.2f})"

    # 4. any qty, name similarity ≥ 0.65
    if pd.notna(comp):
        best = None; best_s = 0
        for idx, r in pool.iterrows():
            s = name_sim(comp, r["Product"])
            if s > best_s:
                best_s, best = s, idx
        if best_s >= 0.65 and best is not None:
            return best, f"name({best_s:.2f})"

    return None, "unmatched"


# Run matcher, consume iQuote rows as we go
iq_pool = iq.copy()
iq_pool["_consumed"] = False
matches = []
methods = []
for _, s in spf.iterrows():
    available = iq_pool[~iq_pool["_consumed"]]
    idx, method = match_iq(s, available)
    if idx is not None:
        iq_pool.loc[idx, "_consumed"] = True
    matches.append(idx)
    methods.append(method)

spf["_iq_idx"] = matches
spf["Match Method"] = methods

# Merge iQuote fields onto SPOT rows
iq_view = iq.rename(columns={"Product": "iq_product", "Quantity": "iq_qty",
                             "ComponentTotalPrice": "iq_component_total",
                             "OptDesc": "iq_optdesc"})
m = spf.merge(
    iq_view[["iq_product","iq_qty","iq_component_total","iq_optdesc"]],
    left_on="_iq_idx", right_index=True, how="left",
)


# Backfill customer from Round 2 SPOT export
cust_map = r2.dropna(subset=["SPOT Estimate #"]).drop_duplicates("SPOT Estimate #").set_index("SPOT Estimate #")["Customer"].to_dict()
m["Customer"] = m["SPOT Estimate #"].map(cust_map)


# Job type
def num(x):
    try: v = float(x); return 0.0 if pd.isna(v) else v
    except (TypeError, ValueError): return 0.0
def has_print(r):
    pm = r.get("Print Mode")
    if pd.notna(pm) and str(pm).strip().lower() not in ("","none","no","nan"): return True
    return num(r.get("Print Cost")) > 0
def has_finishing(r):
    return (num(r.get("Strip Cost")) + num(r.get("Glue Cost")) + num(r.get("Tape Cost"))) > 0
def job_type(r):
    if has_finishing(r): return "Print + Cut + Finishing"
    if has_print(r):     return "Print + Cut"
    return "Cut Only"
m["Job Type"] = m.apply(job_type, axis=1)


m["SPOT Component Total"] = pd.to_numeric(m["Sell Total"], errors="coerce")
m["iQuote Component Total"] = pd.to_numeric(m["iq_component_total"], errors="coerce")
m["Variance $"] = m["SPOT Component Total"] - m["iQuote Component Total"]
m["Variance %"] = np.where(
    m["iQuote Component Total"].fillna(0) != 0,
    m["Variance $"] / m["iQuote Component Total"] * 100.0,
    np.nan,
)


def qty_band(q):
    if pd.isna(q): return "Unknown"
    q = int(q)
    if q < 50:      return "1-49"
    if q < 200:     return "50-199"
    if q < 1000:    return "200-999"
    if q < 5000:    return "1000-4999"
    return "5000+"

m["Qty Band"] = m["Qty"].apply(qty_band)

BANDS = ["1-49","50-199","200-999","1000-4999","5000+"]
JOB_TYPES = ["Cut Only","Print + Cut","Print + Cut + Finishing"]


# ---------- Aggregate ----------
def stats_by_band(df):
    rows = []
    for b in BANDS:
        s = df.loc[df["Qty Band"] == b, "Variance %"].dropna()
        rows.append([b, len(s),
                     s.mean() if len(s) else np.nan,
                     s.median() if len(s) else np.nan])
    return pd.DataFrame(rows, columns=["Qty Band","Count","Avg Variance %","Median Variance %"])

by_type = {jt: stats_by_band(m[m["Job Type"] == jt]) for jt in JOB_TYPES}


# ---------- Workbook ----------
wb = Workbook()
ws = wb.active
ws.title = "Job-Type Variance"

TITLE = Font(bold=True, size=16, color="1F4E78")
SUBHDR = Font(bold=True, size=12, color="1F4E78")
SUBTITLE = Font(italic=True, size=10, color="595959")
HDR = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
BOLD = Font(bold=True)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

FMT_PCT = "0.00\"%\";-0.00\"%\";0.00\"%\""
FMT_INT = "#,##0"
FMT_MONEY = "$#,##0.00;[Red]-$#,##0.00"

ws["A1"] = "SPOT vs iQuote — Variance by Job Type × Qty Band  (v2, corrected join)"
ws["A1"].font = TITLE
ws["A2"] = (
    "Component-level analysis. v2 replaces the ordinal fallback with a fuzzy name+qty matcher. "
    "Rows that could not be matched with confidence are excluded from stats (not forced onto arbitrary partners)."
)
ws["A2"].font = SUBTITLE
ws["A3"] = (
    f"Filter: only the {len(r2_ids)} estimates in the Round 2 quote-total analysis. "
    "-COPY duplicates dropped. Match method is on each row of the list sheets for audit."
)
ws["A3"].font = SUBTITLE
ws.merge_cells("A1:P1"); ws.merge_cells("A2:P2"); ws.merge_cells("A3:P3")


# Match-method breakdown block
mm_start = 5
ws.cell(row=mm_start, column=1, value="Match Method Breakdown (all rows)").font = SUBHDR
ws.cell(row=mm_start+1, column=1, value="Method").font = BOLD
ws.cell(row=mm_start+1, column=2, value="Count").font = BOLD
mm_counts = m["Match Method"].value_counts()
for i, (mth, cnt) in enumerate(mm_counts.items()):
    ws.cell(row=mm_start+2+i, column=1, value=str(mth))
    ws.cell(row=mm_start+2+i, column=2, value=int(cnt)).number_format = FMT_INT


# Job-Type totals table
totals_start = mm_start + 2 + len(mm_counts) + 2
ws.cell(row=totals_start, column=1, value="Component Counts by Job Type").font = SUBHDR
hdr_row = totals_start + 1
for j, h in enumerate(["Job Type","Components","Avg Variance %","Median Variance %"], start=1):
    c = ws.cell(row=hdr_row, column=j, value=h)
    c.font = HDR; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER

for i, jt in enumerate(JOB_TYPES):
    r = hdr_row + 1 + i
    df = m[m["Job Type"] == jt]
    vs = df["Variance %"].dropna()
    ws.cell(row=r, column=1, value=jt).font = BOLD
    ws.cell(row=r, column=2, value=int(len(vs))).number_format = FMT_INT
    ws.cell(row=r, column=3, value=float(vs.mean()) if len(vs) else None).number_format = FMT_PCT
    ws.cell(row=r, column=4, value=float(vs.median()) if len(vs) else None).number_format = FMT_PCT
    for cc in range(1, 5):
        ws.cell(row=r, column=cc).border = BORDER


# Per-band tables + charts
row = totals_start + len(JOB_TYPES) + 4
tbl_meta = []
for jt in JOB_TYPES:
    ws.cell(row=row, column=1, value=jt).font = SUBHDR
    row += 1
    for j, h in enumerate(["Qty Band","Count","Avg Variance %","Median Variance %"], start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = HDR; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    hdr = row
    row += 1
    start = row
    df = by_type[jt]
    for _, dr in df.iterrows():
        ws.cell(row=row, column=1, value=dr["Qty Band"])
        c = ws.cell(row=row, column=2, value=int(dr["Count"])); c.number_format = FMT_INT
        v = dr["Avg Variance %"]
        c = ws.cell(row=row, column=3, value=None if pd.isna(v) else float(v)); c.number_format = FMT_PCT
        v = dr["Median Variance %"]
        c = ws.cell(row=row, column=4, value=None if pd.isna(v) else float(v)); c.number_format = FMT_PCT
        for cc in range(1, 5): ws.cell(row=row, column=cc).border = BORDER
        row += 1
    tbl_meta.append((jt, hdr, start, start + len(df) - 1))
    row += 1

# Charts
anchors = [f"F{tbl_meta[0][1]-1}", f"F{tbl_meta[1][1]-1}", f"F{tbl_meta[2][1]-1}"]
for (jt, hdr, start, end), anchor in zip(tbl_meta, anchors):
    ch = BarChart(); ch.type = "col"; ch.style = 11
    ch.title = f"{jt} — Avg vs Median Variance % by Qty Band"
    ch.y_axis.title = "Variance %"; ch.x_axis.title = "Qty Band"
    ch.height = 8; ch.width = 22
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=hdr, max_row=end), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=start, max_row=end))
    ws.add_chart(ch, anchor)


# Combined chart
comb_top = row + 2
ws.cell(row=comb_top, column=1, value="Combined — Median Variance % by Job Type × Qty Band").font = SUBHDR
comb_hdr = comb_top + 1
ws.cell(row=comb_hdr, column=1, value="Qty Band")
for j, jt in enumerate(JOB_TYPES, start=2):
    ws.cell(row=comb_hdr, column=j, value=jt)
for c in range(1, 5):
    cell = ws.cell(row=comb_hdr, column=c); cell.font = HDR; cell.fill = HDR_FILL
    cell.alignment = CENTER; cell.border = BORDER

for i, b in enumerate(BANDS):
    r = comb_hdr + 1 + i
    ws.cell(row=r, column=1, value=b)
    for j, jt in enumerate(JOB_TYPES, start=2):
        vs = m[(m["Job Type"] == jt) & (m["Qty Band"] == b)]["Variance %"].dropna()
        cell = ws.cell(row=r, column=j, value=float(vs.median()) if len(vs) else None)
        cell.number_format = FMT_PCT
    for cc in range(1, 5): ws.cell(row=r, column=cc).border = BORDER

comb = BarChart(); comb.type = "col"; comb.style = 13
comb.title = "Combined — Median Variance % by Job Type × Qty Band"
comb.y_axis.title = "Median Variance %"; comb.x_axis.title = "Qty Band"
comb.height = 10; comb.width = 24
comb.add_data(Reference(ws, min_col=2, max_col=4, min_row=comb_hdr, max_row=comb_hdr + len(BANDS)),
              titles_from_data=True)
comb.set_categories(Reference(ws, min_col=1, min_row=comb_hdr + 1, max_row=comb_hdr + len(BANDS)))
ws.add_chart(comb, f"F{comb_top}")


ws.column_dimensions["A"].width = 28
for c in "BCD": ws.column_dimensions[c].width = 18
for c in "EFGHIJKLMNOP": ws.column_dimensions[c].width = 12


# Per-job-type lists
band_sort = {"1-49":0,"50-199":1,"200-999":2,"1000-4999":3,"5000+":4,"Unknown":5}
def write_list(name, sub):
    s = wb.create_sheet(name)
    s["A1"] = f"{name} — Component List (v2, corrected join)"
    s["A1"].font = TITLE
    s["A2"] = "Rows are components. Match Method shows how each SPOT row was paired to iQuote."
    s["A2"].font = SUBTITLE
    s.merge_cells("A1:L1"); s.merge_cells("A2:L2")

    headers = ["Qty Band","iQuote Estimate #","SPOT Estimate #","Customer","Project Name",
               "SPOT Component","Qty","iQuote Product","iQuote Qty","Print Mode",
               "SPOT Sell Total","iQuote Total","Variance %","Match Method"]
    for j, h in enumerate(headers, start=1):
        c = s.cell(row=4, column=j, value=h)
        c.font = HDR; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER

    sub = sub.copy()
    sub["_bk"] = sub["Qty Band"].map(band_sort).fillna(99)
    sub = sub.sort_values(["_bk","Customer","iq_id","Qty"]).reset_index(drop=True)

    for i in range(len(sub)):
        r = 5 + i; row_ = sub.iloc[i]
        s.cell(row=r, column=1, value=row_["Qty Band"])
        s.cell(row=r, column=2, value=int(row_["iq_id"]) if pd.notna(row_["iq_id"]) else None).number_format = FMT_INT
        s.cell(row=r, column=3, value=row_["SPOT Estimate #"])
        s.cell(row=r, column=4, value=row_["Customer"])
        s.cell(row=r, column=5, value=row_["Project Name"])
        s.cell(row=r, column=6, value=row_["Component"])
        s.cell(row=r, column=7, value=int(row_["Qty"]) if pd.notna(row_["Qty"]) else None).number_format = FMT_INT
        s.cell(row=r, column=8, value=row_.get("iq_product"))
        s.cell(row=r, column=9, value=int(row_["iq_qty"]) if pd.notna(row_.get("iq_qty")) else None).number_format = FMT_INT
        pm = row_["Print Mode"]
        s.cell(row=r, column=10, value=None if pd.isna(pm) else str(pm))
        v = row_["SPOT Component Total"]
        s.cell(row=r, column=11, value=float(v) if pd.notna(v) else None).number_format = FMT_MONEY
        v = row_["iQuote Component Total"]
        s.cell(row=r, column=12, value=float(v) if pd.notna(v) else None).number_format = FMT_MONEY
        v = row_["Variance %"]
        s.cell(row=r, column=13, value=float(v) if pd.notna(v) else None).number_format = FMT_PCT
        s.cell(row=r, column=14, value=str(row_["Match Method"]))
        for cc in range(1, 15): s.cell(row=r, column=cc).border = BORDER

    s.auto_filter.ref = f"A4:N{4 + len(sub)}"
    s.freeze_panes = "A5"
    widths = [12,18,22,26,38,26,10,26,10,14,16,16,14,20]
    for j, w in enumerate(widths, start=1):
        s.column_dimensions[get_column_letter(j)].width = w


list_cols = m[["Qty Band","iq_id","SPOT Estimate #","Customer","Project Name",
               "Component","Qty","iq_product","iq_qty","Print Mode","SPOT Component Total",
               "iQuote Component Total","Variance %","Match Method","Job Type"]]
sheet_map = {"Cut Only":"Cut Only List","Print + Cut":"Print + Cut List",
             "Print + Cut + Finishing":"Print+Cut+Fin List"}
for jt in JOB_TYPES:
    write_list(sheet_map[jt], list_cols[list_cols["Job Type"] == jt].drop(columns=["Job Type"]))

# Unmatched sheet
unm = list_cols[list_cols["Match Method"].isin(["unmatched","no-iq-rows"])].drop(columns=["Job Type"])
if len(unm) > 0:
    write_list("Unmatched", unm)

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Total components: {len(m)}   With variance: {m['Variance %'].notna().sum()}")
print()
print("Match method breakdown:")
print(mm_counts.to_string())
print()
print("Job-Type summary (v2):")
for jt in JOB_TYPES:
    vs = m[m["Job Type"] == jt]["Variance %"].dropna()
    print(f"  {jt:>28}: n={len(vs):3d}  mean={vs.mean():+6.1f}%  median={vs.median():+6.1f}%")
print()
print("Print + Cut by band (v2):")
for b in BANDS:
    vs = m[(m["Job Type"] == "Print + Cut") & (m["Qty Band"] == b)]["Variance %"].dropna()
    if len(vs):
        print(f"  {b:>10}: n={len(vs):2d}  mean={vs.mean():+6.1f}%  median={vs.median():+6.1f}%")
