"""Job-Type Variance Analysis — single-sheet workbook.

Filters full SPOT dataset to the 110 estimates already in the Round 2 analysis.
Component-level classification into Cut Only / Print+Cut / Print+Cut+Finishing.

Usage: python3 generate_jobtype.py <full_spot.xlsx> <round2_input.xlsx> <output.xlsx>
"""
import sys
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


SPOT_FULL = sys.argv[1] if len(sys.argv) > 1 else "spot_full.xlsx"
R2_INPUT  = sys.argv[2] if len(sys.argv) > 2 else "input.xlsx"
OUT       = sys.argv[3] if len(sys.argv) > 3 else "SPOT_JobType_Variance.xlsx"


# ---------- Load ----------
spf = pd.read_excel(SPOT_FULL, sheet_name="Testing Comparison")
iq  = pd.read_excel(R2_INPUT,  sheet_name="iQuote Data")
r2  = pd.read_excel(R2_INPUT,  sheet_name="SPOT Data")

r2_ids = set(r2["iQuote Estimate #"].dropna().astype(int).unique())

spf["iq_id"] = pd.to_numeric(spf["iQuote #"], errors="coerce").astype("Int64")

# Drop -COPY duplicates and filter to Round 2 estimates only
spf = spf[~spf["SPOT Estimate #"].astype(str).str.contains("-COPY", na=False)]
spf = spf[spf["iq_id"].isin(r2_ids)].reset_index(drop=True)


def norm_opt(x):
    if pd.isna(x): return ""
    s = str(x).strip().lower()
    if " - " in s: s = s.split(" - ")[0].strip()
    return s


spf["_opt"] = spf["iQuote Option"].apply(norm_opt)
iq["_opt"]  = iq["OptDesc"].apply(norm_opt)


# ---------- Job-Type classification ----------
def num(x):
    try:
        v = float(x)
        return 0.0 if pd.isna(v) else v
    except (TypeError, ValueError):
        return 0.0


def has_print(row):
    pm = row.get("Print Mode")
    if pd.notna(pm) and str(pm).strip().lower() not in ("", "none", "no", "nan"):
        return True
    return num(row.get("Print Cost")) > 0


def has_finishing(row):
    return (num(row.get("Strip Cost"))
            + num(row.get("Glue Cost"))
            + num(row.get("Tape Cost"))) > 0


def job_type(row):
    if has_finishing(row): return "Print + Cut + Finishing"
    if has_print(row):     return "Print + Cut"
    return "Cut Only"


spf["Job Type"] = spf.apply(job_type, axis=1)


# ---------- Join to iQuote (component-level, same rule as Round 2) ----------
spf_sorted = spf.sort_values(["iq_id","_opt","Qty"]).reset_index(drop=True)
iq_sorted  = iq.sort_values(["IdQuote","_opt","Quantity"]).reset_index(drop=True)
spf_sorted["_sp_idx"] = spf_sorted.index
iq_sorted["_iq_idx"]  = iq_sorted.index

pool = iq_sorted.copy()
pool["_consumed"] = False
pairs = []
for _, s in spf_sorted.iterrows():
    p = pool[(pool["IdQuote"] == s["iq_id"]) & (pool["_opt"] == s["_opt"])
             & (pool["Quantity"] == s["Qty"]) & (~pool["_consumed"])]
    if len(p) == 0:
        p = pool[(pool["IdQuote"] == s["iq_id"]) & (pool["Quantity"] == s["Qty"])
                 & (~pool["_consumed"])]
    if len(p) == 0:
        p = pool[(pool["IdQuote"] == s["iq_id"]) & (~pool["_consumed"])]
    if len(p) > 0:
        pick = p.iloc[0]
        pairs.append((s["_sp_idx"], pick["_iq_idx"]))
        pool.loc[pick["_iq_idx"], "_consumed"] = True
    else:
        pairs.append((s["_sp_idx"], None))

mdf = pd.DataFrame(pairs, columns=["_sp_idx", "_iq_idx"])
m = spf_sorted.merge(mdf, on="_sp_idx").merge(
    iq_sorted.drop(columns=["_iq_idx", "_opt"]),
    left_on="_iq_idx", right_index=True, how="left"
)

m["SPOT Component Total"] = pd.to_numeric(m["Sell Total"], errors="coerce")
m["iQuote Component Total"] = pd.to_numeric(m["ComponentTotalPrice"], errors="coerce")
m["Variance $"] = m["SPOT Component Total"] - m["iQuote Component Total"]
m["Variance %"] = np.where(
    m["iQuote Component Total"].fillna(0) != 0,
    m["Variance $"] / m["iQuote Component Total"] * 100.0,
    np.nan,
)

# Qty band
def qty_band(q):
    if pd.isna(q): return "Unknown"
    q = int(q)
    if q < 50:      return "1-49"
    if q < 200:     return "50-199"
    if q < 1000:    return "200-999"
    if q < 5000:    return "1000-4999"
    return "5000+"

m["Qty Band"] = m["Qty"].apply(qty_band)

BANDS = ["1-49", "50-199", "200-999", "1000-4999", "5000+"]
JOB_TYPES = ["Cut Only", "Print + Cut", "Print + Cut + Finishing"]


# ---------- Aggregate ----------
def stats_by_band(df):
    rows = []
    for b in BANDS:
        s = df.loc[df["Qty Band"] == b, "Variance %"].dropna()
        rows.append([
            b,
            len(s),
            s.mean() if len(s) else np.nan,
            s.median() if len(s) else np.nan,
        ])
    return pd.DataFrame(rows, columns=["Qty Band", "Count", "Avg Variance %", "Median Variance %"])


by_type = {jt: stats_by_band(m[m["Job Type"] == jt]) for jt in JOB_TYPES}


# ---------- Write workbook ----------
wb = Workbook()
ws = wb.active
ws.title = "Job-Type Variance"

TITLE   = Font(bold=True, size=16, color="1F4E78")
SUBHDR  = Font(bold=True, size=12, color="1F4E78")
SUBTITLE = Font(italic=True, size=10, color="595959")
HDR     = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL= PatternFill("solid", fgColor="1F4E78")
BOLD    = Font(bold=True)
THIN    = Side(border_style="thin", color="BFBFBF")
BORDER  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER  = Alignment(horizontal="center", vertical="center")

FMT_PCT   = "0.00\"%\";-0.00\"%\";0.00\"%\""
FMT_INT   = "#,##0"

ws["A1"] = "SPOT vs iQuote — Variance by Job Type × Qty Band"
ws["A1"].font = TITLE
ws["A2"] = "Component-level analysis. Job Type is determined per component from Print/Strip/Glue/Tape cost signals."
ws["A2"].font = SUBTITLE
ws["A3"] = (
    f"Filter: only the {len(r2_ids)} estimates used in the Round 2 quote-total analysis. "
    "Any additional estimates in the SPOT export are excluded. -COPY duplicates dropped."
)
ws["A3"].font = SUBTITLE
ws.merge_cells("A1:P1"); ws.merge_cells("A2:P2"); ws.merge_cells("A3:P3")

# Job-Type totals table (top-left)
ws["A5"] = "Component Counts by Job Type"; ws["A5"].font = SUBHDR
ws["A6"] = "Job Type"; ws["B6"] = "Components"; ws["C6"] = "Avg Variance %"; ws["D6"] = "Median Variance %"
for c in range(1, 5):
    cell = ws.cell(row=6, column=c); cell.font = HDR; cell.fill = HDR_FILL
    cell.alignment = CENTER; cell.border = BORDER

for i, jt in enumerate(JOB_TYPES):
    r = 7 + i
    df = m[m["Job Type"] == jt]
    vs = df["Variance %"].dropna()
    ws.cell(row=r, column=1, value=jt).font = BOLD
    cc = ws.cell(row=r, column=2, value=len(df)); cc.number_format = FMT_INT
    ma = ws.cell(row=r, column=3, value=float(vs.mean()) if len(vs) else None); ma.number_format = FMT_PCT
    md = ws.cell(row=r, column=4, value=float(vs.median()) if len(vs) else None); md.number_format = FMT_PCT
    for c in range(1, 5):
        ws.cell(row=r, column=c).border = BORDER

# Per-band tables — one per job type stacked vertically below KPIs
row = 12
tbl_starts = []
for jt in JOB_TYPES:
    ws.cell(row=row, column=1, value=jt).font = SUBHDR
    row += 1
    headers = ["Qty Band", "Count", "Avg Variance %", "Median Variance %"]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=j, value=h)
        cell.font = HDR; cell.fill = HDR_FILL; cell.alignment = CENTER; cell.border = BORDER
    tbl_header_row = row
    row += 1
    tbl_start_row = row
    df = by_type[jt]
    for _, dr in df.iterrows():
        ws.cell(row=row, column=1, value=dr["Qty Band"])
        c = ws.cell(row=row, column=2, value=int(dr["Count"])); c.number_format = FMT_INT
        c = ws.cell(row=row, column=3, value=None if pd.isna(dr["Avg Variance %"]) else float(dr["Avg Variance %"])); c.number_format = FMT_PCT
        c = ws.cell(row=row, column=4, value=None if pd.isna(dr["Median Variance %"]) else float(dr["Median Variance %"])); c.number_format = FMT_PCT
        for cc in range(1, 5):
            ws.cell(row=row, column=cc).border = BORDER
        row += 1
    tbl_starts.append((jt, tbl_header_row, tbl_start_row, tbl_start_row + len(df) - 1))
    row += 1  # blank spacer

# ---------- Charts — one grouped bar per job type ----------
chart_anchors = ["F5", "F22", "F39"]
for (jt, hdr_row, start_row, end_row), anchor in zip(tbl_starts, chart_anchors):
    ch = BarChart()
    ch.type = "col"
    ch.style = 11
    ch.title = f"{jt} — Avg vs Median Variance % by Qty Band"
    ch.y_axis.title = "Variance %"
    ch.x_axis.title = "Qty Band"
    ch.height = 8
    ch.width = 22
    data_ref = Reference(ws, min_col=3, max_col=4, min_row=hdr_row, max_row=end_row)
    cats_ref = Reference(ws, min_col=1, min_row=start_row, max_row=end_row)
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ws.add_chart(ch, anchor)


# Combined chart at bottom — all 3 job types on same axis
# Build a small pivot table for the combined chart
combined_start_row = row + 2
ws.cell(row=combined_start_row, column=1, value="Combined — Median Variance % by Job Type × Qty Band").font = SUBHDR
combined_hdr = combined_start_row + 1
ws.cell(row=combined_hdr, column=1, value="Qty Band")
for j, jt in enumerate(JOB_TYPES, start=2):
    ws.cell(row=combined_hdr, column=j, value=jt)
for c in range(1, 5):
    cell = ws.cell(row=combined_hdr, column=c); cell.font = HDR; cell.fill = HDR_FILL
    cell.alignment = CENTER; cell.border = BORDER

for i, b in enumerate(BANDS):
    r = combined_hdr + 1 + i
    ws.cell(row=r, column=1, value=b)
    for j, jt in enumerate(JOB_TYPES, start=2):
        vs = m[(m["Job Type"] == jt) & (m["Qty Band"] == b)]["Variance %"].dropna()
        cell = ws.cell(row=r, column=j, value=float(vs.median()) if len(vs) else None)
        cell.number_format = FMT_PCT
    for cc in range(1, 5):
        ws.cell(row=r, column=cc).border = BORDER

comb = BarChart(); comb.type = "col"; comb.style = 13
comb.title = "Combined — Median Variance % by Job Type × Qty Band"
comb.y_axis.title = "Median Variance %"
comb.x_axis.title = "Qty Band"
comb.height = 10; comb.width = 24
comb.add_data(Reference(ws, min_col=2, max_col=4, min_row=combined_hdr, max_row=combined_hdr + len(BANDS)),
              titles_from_data=True)
comb.set_categories(Reference(ws, min_col=1, min_row=combined_hdr + 1, max_row=combined_hdr + len(BANDS)))
ws.add_chart(comb, f"F{combined_start_row}")


# Column widths
ws.column_dimensions["A"].width = 28
for c in "BCD":
    ws.column_dimensions[c].width = 18
for c in "EFGHIJKLMNOP":
    ws.column_dimensions[c].width = 12

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Components analyzed: {len(m)}   Estimates: {m['iq_id'].nunique()}")
print("Job-Type breakdown:")
print(m['Job Type'].value_counts())
print()
print("Summary per Job Type:")
for jt in JOB_TYPES:
    vs = m[m['Job Type']==jt]['Variance %'].dropna()
    print(f"  {jt:>28}: n={len(vs):4d}   mean={vs.mean():+.1f}%   median={vs.median():+.1f}%")
