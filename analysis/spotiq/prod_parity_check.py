"""READ-ONLY parity check against PRODUCTION: prod report xlsx vs prod engine
(UI-equivalent payloads) for every flagged quote. No writes to prod."""
import io
import sys

import openpyxl
import requests

sys.path.insert(0, '/app/analysis/spotiq')
from parity_regression import build_mp_payload, build_payload, js_state, r2

PROD = "https://ai-connect-33.emergent.host"
TOKEN = requests.post(f"{PROD}/api/auth/login", json={
    "email": "admin@standardprinting.com", "password": "spot-admin-2026"}).json()["access_token"]
H = {"Authorization": f"Bearer {TOKEN}"}

quotes = requests.get(f"{PROD}/api/quotes?scope=team&limit=500", headers=H).json()
flagged = [q for q in quotes if (q.get("state") or {}).get("goodForTesting")]
print(f"prod: {len(quotes)} quotes, {len(flagged)} flagged for testing")


def ui_values(q):
    state = js_state(q.get("state") or {})
    is_mp = state.get("quoteType") == "Multi-Part" and bool(state.get("bomItems"))
    out = {}
    if is_mp:
        data = requests.post(f"{PROD}/api/engine/calculate-multi-part", headers=H,
                             json=build_mp_payload(state, q.get("id"))).json()
        per_job = state.get("freightScope") == "Per Job"
        for slot in (data.get("slots") or []):
            for pc in (slot.get("perComponent") or []):
                if pc.get("error"):
                    continue
                qt = r2(slot.get("sellTotal")) if per_job else r2(pc.get("sellTotal"))
                fr = r2(pc.get("allocatedFreight")) if per_job else r2(
                    (pc.get("freightLine") or 0) + (pc.get("freightEmbed") or 0))
                out[(pc.get("desc") or "", pc.get("qty"))] = (
                    r2(pc.get("itemSell")), qt, r2(pc.get("totalMfg")),
                    r2(pc.get("gmDollar")), round(float(pc.get("gmPct") or 0) * 100, 1), fr)
    else:
        data = requests.post(f"{PROD}/api/engine/calculate-all", headers=H,
                             json=build_payload(state, q.get("id"))).json()
        desc = state.get("itemDesc") or "Item"
        for rr in (data.get("results") or []):
            if rr.get("error"):
                continue
            out[(desc, rr.get("qty"))] = (
                r2(rr.get("itemSell")), r2(rr.get("sellTotal")), r2(rr.get("totalMfg")),
                r2(rr.get("gmDollar")), round(float(rr.get("gmPct") or 0) * 100, 1),
                r2((rr.get("freightLine") or 0) + (rr.get("freightEmbed") or 0)))
    return out


r = requests.get(f"{PROD}/api/quotes/_report/testing.xlsx", headers=H)
wb = openpyxl.load_workbook(io.BytesIO(r.content))
ws = wb.active
hdr = [c.value for c in ws[1]]
rep = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(hdr, row))
    comp = d.get("Component") or ""
    if str(comp).startswith("ENGINE ERROR"):
        print("REPORT ERROR ROW:", d.get("SPOT Estimate #"), comp[:120])
        continue
    rep[(d.get("SPOT Estimate #"), comp, d.get("Qty"))] = (
        d.get("Sell Total"), d.get("Quote Total"), d.get("Total Cost"),
        d.get("Margin $"), d.get("Margin %"),
        d.get("Freight") if d.get("Freight") != "" else 0.0)

n_pass = n_fail = 0
for q in flagged:
    est = q.get("est_num")
    try:
        ui = ui_values(q)
    except Exception as e:  # noqa: BLE001
        print(f"ERROR {est}: {type(e).__name__}: {e}")
        n_fail += 1
        continue
    issues = []
    for (comp, qty), exp in sorted(ui.items(), key=lambda x: (str(x[0][0]), str(x[0][1]))):
        got = rep.get((est, comp, qty))
        if got != exp:
            issues.append(f"   {comp}@{qty}: UI(sell,qt,cost,gm$,gm%,fr)={exp} REPORT={got}")
    if issues:
        n_fail += 1
        print(f"FAIL {est} [{(q.get('state') or {}).get('quoteType')}/"
              f"{(q.get('state') or {}).get('pricingMode', '')}]")
        for i in issues[:8]:
            print(i)
    else:
        n_pass += 1
        print(f"PASS {est}")

print(f"\n=== PROD PARITY: {n_pass} PASS / {n_fail} FAIL of {len(flagged)} flagged quotes ===")
