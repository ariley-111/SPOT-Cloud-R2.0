import re, statistics
import openpyxl
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

import sys
sys.path.insert(0, '/app/analysis/spotiq')
from matcher import matches, JOBS, spot_cat, iq_cat, act_cat, f

for m in matches:
    m['dpct'] = (m['spot_sell'] - m['iq_sell']) / m['iq_sell'] * 100 if m['iq_sell'] else 0
    m['verdict'] = 'PASS' if abs(m['dpct']) <= 10 else ('REVIEW' if abs(m['dpct']) <= 20 else 'FAIL')
OUT = {('RW-260723-084853', '8'), ('RW-260723-084853', '15'), ('RW-260720-082503', '918+1750')}
for m in matches:
    m['outlier'] = (m['est'], m['qty']) in OUT

STOP = {'print', 'printed', 'digital', 'plain', 'white', 'kraft', 'scored', 'glued', 'stripped',
        'outside', 'inside', 'only', 'and', 'with', 'no', '1c', '2c', 'direct', 'cut', 'cad',
        'coverage', '100', 'sides', 'the', 'a', 'for', 'x', 'set', 'sided', 'double'}


def toks(s):
    return {t for t in re.findall(r'[a-z0-9#/\-]+', str(s).lower()) if len(t) > 1 and t not in STOP}


def match_components(s_rows, i_rows):
    def qm2(a, b):
        try:
            a, b = float(a), float(b)
        except (TypeError, ValueError):
            return False
        return a == b or abs(a - b) <= max(2, 0.02 * max(a, b))
    if len(s_rows) == 1 and len(i_rows) == 1:
        return {id(s_rows[0]): i_rows[0]}
    pairs = []
    for sr in s_rows:
        st = toks(sr['Component'])
        for ir in i_rows:
            it = toks(ir['Product'])
            shared = len(st & it)
            if shared >= 1:
                pairs.append((shared * 2 + (1 if qm2(sr['Qty'], ir['Quantity']) else 0), id(sr), sr, ir))
    pairs.sort(key=lambda p: -p[0])
    used_s, used_i, out = set(), set(), {}
    for score, sid, sr, ir in pairs:
        if sid in used_s or id(ir) in used_i:
            continue
        used_s.add(sid)
        used_i.add(id(ir))
        out[sid] = ir
    return out


def spec_text(r):
    bits = []
    pm = r['Print Mode']
    if pm and f(r['Print Cost']) > 0:
        bits.append(f"{pm} print")
    elif f(r['Print Cost']) == 0:
        bits.append("No print")
    route = r['Route'] or ''
    if route == 'CAD':
        bits.append("CAD Cut")
    elif route == 'Crest':
        bits.append("Die Cut (Crest)")
    elif route == 'Bobst':
        bits.append("Die Cut (Bobst)")
    elif route:
        bits.append(route)
    if f(r['Strip Cost']) > 0:
        bits.append("Hand Stripped")
    if f(r['Glue Cost']) > 0:
        gj = r['Glue Joints']
        bits.append(f"Glued{f' ({gj} jt)' if gj else ''}")
    if r['Tape Type']:
        bits.append(f"Tape ({r['Tape Type']} x{r['Tape Strip Count']})")
    ir_ = r['Inches of Rule']
    if ir_:
        bits.append(f"{f(ir_):.0f}\" rule")
    return " · ".join(bits)


GREEN = PatternFill('solid', start_color='D9EAD3')
AMBER = PatternFill('solid', start_color='FFF3CD')
RED = PatternFill('solid', start_color='F4CCCC')
HDRF = PatternFill('solid', start_color='1F3864')
GREYF = PatternFill('solid', start_color='EDEFF5')
HDR_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE = Font(bold=True, size=15, color='1F3864')
SUB = Font(bold=True, size=11, color='1F3864')
VFILL = {'PASS': GREEN, 'REVIEW': AMBER, 'FAIL': RED}

ACT = {(e, str(q)): j for j, e, q, *_ in [(j, e, q) for j, e, q, qid, o, iq2, n in JOBS]}

wb = openpyxl.Workbook()

# ═══ Sheet 1: Executive Summary ═══
ws = wb.active
ws.title = 'Executive Summary'
ws.column_dimensions['A'].width = 60
ws.column_dimensions['B'].width = 26
ws['A1'] = 'SPOT Launch Readiness — Executive Summary'
ws['A1'].font = TITLE
ws['A2'] = 'SPOT (new estimator) vs iQuote (current) on 38 real quotes, validated against actual production costs on 9 produced jobs. Verdicts use legacy parity thresholds: PASS ≤±10%, REVIEW ≤±20%, FAIL >±20%.'
ws['A2'].font = Font(size=9, italic=True)
ws['A2'].alignment = Alignment(wrap_text=True)
ws.merge_cells('A2:B2')
ws.row_dimensions[2].height = 30

clean = [m for m in matches if not m['outlier']]
d = [m['dpct'] for m in clean]
rows = [
    ('HEADLINE', ''),
    ('Comparable pricing scenarios analyzed', f"{len(matches)} (3 outliers isolated)"),
    ('SPOT prices vs iQuote (median)', f"{statistics.median(d):+.1f}%"),
    ('Scenarios within ±10% / ±20%', f"{sum(1 for x in d if abs(x)<=10)} / {sum(1 for x in d if abs(x)<=20)} of {len(d)}"),
    ('', ''),
    ('GROUND TRUTH — 9 PRODUCED JOBS', ''),
    ('Invoiced vs actual cost', '$18,131 invoiced / $22,305 cost = -23% margin'),
    ('Jobs that lost money as invoiced', '5 of 9'),
    ('Same work at SPOT prices', '$21,277 = -5% (near break-even)'),
    ('', ''),
    ('ACCURACY vs ACTUAL COST (per category)', ''),
    ('Material', 'SPOT within ±10% — well calibrated'),
    ('Print', 'SPOT -14% vs iQuote -49% — iQuote badly under-costs print'),
    ('Cutting', 'Both under: SPOT -32%, iQuote -41%. Root cause found (see detailed workbook): assumed cut speed 650 in/min (inherited from legacy) vs ~165 real. Fix identified, BOOKMARKED pending leadership review'),
    ('Freight', 'Within ±7%'),
    ('', ''),
    ('KEY TAKEAWAYS', ''),
    ('1.', "iQuote has been winning work below cost. SPOT's ~13% higher pricing is largely justified by real production costs, not over-estimation."),
    ('2.', 'SPOT is closer to actual cost than iQuote in every measurable category on produced jobs.'),
    ('3.', 'Cutting throughput is the one systematic gap — in BOTH systems, inherited from legacy. Fix is two admin-console values in SPOT (no code); recommendation bookmarked.'),
    ('4.', 'Small quantities (≤50) carry a heavy SPOT premium (+45-88%) — review minimums policy; may be pricing truth, may cost wins.'),
    ('5.', 'Two items need investigation before full trust: eggcrate component anomaly (RW-260723-084853, ~18x iQuote) and one kit modeling mismatch (iQ 223583, won at ~2x SPOT).'),
]
r = 4
for a, b in rows:
    ca = ws.cell(row=r, column=1, value=a)
    cb = ws.cell(row=r, column=2, value=b)
    ca.font = SUB if (a and a == a.upper() and len(a) > 3 and not a[0].isdigit()) else Font(size=10, bold=a in ('1.', '2.', '3.', '4.', '5.'))
    cb.font = Font(size=10)
    ca.alignment = Alignment(wrap_text=True, vertical='top')
    cb.alignment = Alignment(wrap_text=True, vertical='top')
    if a.startswith(('1.', '2.', '3.', '4.', '5.')) or a in ('Cutting',):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 28
    r += 1

# ═══ Sheet 2: Quote Comparison ═══
ws = wb.create_sheet('Quote Comparison')
widths = [44, 46, 8, 11, 12, 12, 10, 8, 9, 30]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws['A1'] = 'Quote-by-Quote: SPOT vs iQuote vs Actual'
ws['A1'].font = TITLE
ws['A2'] = 'Grouped by SPOT project. All $ are SELL prices. Component-level iQuote prices auto-matched by product name + qty (blank = no confident 1:1 match; totals always compare). Δ% = SPOT vs iQuote (+ = SPOT higher).'
ws['A2'].font = Font(size=9, italic=True)
ws.merge_cells('A2:J2')

hdr = ['Project / Component', 'Specs', 'Qty', 'SPOT Sell $', 'SPOT Sell $/ea', 'iQuote Sell $', 'Δ $', 'Δ %', 'Verdict', 'Actual (produced jobs)']
r = 4
for c, h in enumerate(hdr, 1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.fill = HDRF
    cell.font = HDR_FONT
ws.freeze_panes = 'A5'
r += 1

ACTMAP = {(e, str(q)): (j, qid, o, iq2) for j, e, q, qid, o, iq2, n in JOBS}
order = sorted(matches, key=lambda m: (m['est'], str(m['qty'])))
for m in order:
    rows_ = m['s_rows']
    comp_map = match_components(rows_, m['i_rows'])
    head = f"{m['proj'] or m['est']}:   SPOT {m['est']}  ·  iQuote #{m['iq']}" + ("  ·  BECAME A JOB" if m['job'] == 'Yes' else '')
    if m['outlier']:
        head += "   [OUTLIER — see detailed workbook]"
    cell = ws.cell(row=r, column=1, value=head)
    cell.font = Font(bold=True, size=10, color='1F3864')
    for c in range(1, 11):
        ws.cell(row=r, column=c).fill = GREYF
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    r += 1
    for cr in rows_:
        qty = cr['Qty']
        sell = f(cr['Sell Total'])
        ws.cell(row=r, column=1, value=f"    {cr['Component']}").font = Font(size=9)
        ws.cell(row=r, column=2, value=spec_text(cr)).font = Font(size=9)
        ws.cell(row=r, column=3, value=qty).font = Font(size=9)
        ws.cell(row=r, column=4, value=round(sell, 2)).font = Font(size=9)
        ws.cell(row=r, column=5, value=round(sell / qty, 3) if qty else '').font = Font(size=9)
        ir = comp_map.get(id(cr))
        if ir is not None:
            iq_c = f(ir['ComponentTotalPrice'])
            ws.cell(row=r, column=6, value=round(iq_c, 2)).font = Font(size=9)
            ws.cell(row=r, column=7, value=round(sell - iq_c, 2)).font = Font(size=9)
            if iq_c:
                ws.cell(row=r, column=8, value=f"{(sell-iq_c)/iq_c*100:+.1f}%").font = Font(size=9)
        r += 1
    ws.cell(row=r, column=1, value=f"    TOTAL @ qty {m['qty']}").font = Font(bold=True, size=9)
    ws.cell(row=r, column=4, value=round(m['spot_sell'], 2)).font = Font(bold=True, size=9)
    ws.cell(row=r, column=6, value=round(m['iq_sell'], 2)).font = Font(bold=True, size=9)
    ws.cell(row=r, column=7, value=round(m['spot_sell'] - m['iq_sell'], 2)).font = Font(bold=True, size=9)
    ws.cell(row=r, column=8, value=f"{m['dpct']:+.1f}%").font = Font(bold=True, size=9)
    vc = ws.cell(row=r, column=9, value=m['verdict'])
    vc.font = Font(bold=True, size=9)
    vc.fill = VFILL[m['verdict']]
    ak = ACTMAP.get((m['est'], str(m['qty'])))
    if ak:
        job_id = ak[0]
        a = act_cat(job_id)
        marg = (a['Sell'] - a['Total']) / a['Sell'] * 100
        cell = ws.cell(row=r, column=10, value=f"Job {job_id}: invoiced ${a['Sell']:,.0f}, actual cost ${a['Total']:,.0f} → margin {marg:+.0f}%")
        cell.font = Font(size=9, bold=True, color=('9C0006' if marg < 0 else '2E7D32'))
    if m['note']:
        r += 1
        ws.cell(row=r, column=1, value=f"    note: {m['note']}").font = Font(size=8, italic=True, color='777777')
    r += 2

# ═══ Sheet 3: Charts ═══
ws = wb.create_sheet('Charts')
ws['A1'] = 'Price Comparison — 9 produced jobs'
ws['A1'].font = TITLE
r0 = 4
for c, h in enumerate(['MIS Job', 'Actual Invoice $', 'iQuote Sell $', 'SPOT Sell $'], 1):
    ws.cell(row=r0 - 1, column=c, value=h).font = Font(bold=True)
i = 0
for job_id, est, sq, qid, opt, iqq, note in JOBS:
    s, iq_, a = spot_cat(est, sq), iq_cat(qid, opt, iqq), act_cat(job_id)
    ws.cell(row=r0 + i, column=1, value=str(job_id))
    ws.cell(row=r0 + i, column=2, value=round(a['Sell'], 0))
    ws.cell(row=r0 + i, column=3, value=round(iq_['Sell'], 0))
    ws.cell(row=r0 + i, column=4, value=round(s['Sell'], 0))
    i += 1
ch = BarChart()
ch.title = 'Actual invoice vs iQuote sell vs SPOT sell'
ch.height, ch.width = 10, 22
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=r0 - 1, max_row=r0 + i - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r0 + i - 1))
ws.add_chart(ch, 'F3')

r1 = r0 + i + 3
for c, h in enumerate(['Category', 'Actual $', 'SPOT est $', 'iQuote est $'], 1):
    ws.cell(row=r1 - 1, column=c, value=h).font = Font(bold=True)
agg = {x: defaultdict(float) for x in 'SIA'}
for job_id, est, sq, qid, opt, iqq, note in JOBS:
    if job_id == 236943:
        continue
    s, iq_, a = spot_cat(est, sq), iq_cat(qid, opt, iqq), act_cat(job_id)
    for k in ['Material', 'Print', 'Cutting']:
        agg['S'][k] += s[k]; agg['I'][k] += iq_[k]; agg['A'][k] += a[k]
j = 0
for k in ['Material', 'Print', 'Cutting']:
    ws.cell(row=r1 + j, column=1, value=k)
    ws.cell(row=r1 + j, column=2, value=round(agg['A'][k], 0))
    ws.cell(row=r1 + j, column=3, value=round(agg['S'][k], 0))
    ws.cell(row=r1 + j, column=4, value=round(agg['I'][k], 0))
    j += 1
ch2 = BarChart()
ch2.title = 'Estimated COST accuracy vs actual by category (8 jobs, excl blowout)'
ch2.height, ch2.width = 10, 22
ch2.add_data(Reference(ws, min_col=2, max_col=4, min_row=r1 - 1, max_row=r1 + j - 1), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=r1, max_row=r1 + j - 1))
ws.add_chart(ch2, 'F24')
for col, w in (('A', 12), ('B', 15), ('C', 13), ('D', 12)):
    ws.column_dimensions[col].width = w

wb.save('/app/frontend/public/SPOT_Executive_Summary.xlsx')
print('saved:', wb.sheetnames)
