import openpyxl, statistics
from collections import defaultdict

wb = openpyxl.load_workbook('/app/analysis/spotiq/actuals.xlsx', data_only=True)


def f(x):
    try:
        return float(x or 0)
    except (TypeError, ValueError):
        return 0.0


CATS = ['Material', 'Print', 'Cutting', 'Strip', 'Glue', 'Freight']

ws = wb['Testing Comparison']
hdr = [c.value for c in ws[1]]
spot_rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

ws2 = wb['iQuote']
hdr2 = [c.value for c in ws2[1]]
iq_rows = [dict(zip(hdr2, r)) for r in ws2.iter_rows(min_row=2, values_only=True) if r[0]]

ws3 = wb['Actuals']
hdr3 = [c.value for c in ws3[1]]
act_rows = [dict(zip(hdr3, r)) for r in ws3.iter_rows(min_row=2, values_only=True) if r[0]]


def spot_cats(rows):
    return {'Material': sum(f(r['Material Cost']) for r in rows),
            'Print': sum(f(r['Print Cost']) for r in rows),
            'Cutting': sum(f(r['Cutting Cost']) for r in rows),
            'Strip': sum(f(r['Strip Cost']) for r in rows),
            'Glue': sum(f(r['Glue Cost']) for r in rows),
            'Freight': sum(f(r['Freight']) for r in rows)}


def iq_cats(rows):
    return {'Material': sum(f(r['MaterialCost']) for r in rows),
            'Print': sum(f(r['PressCost']) for r in rows),
            'Cutting': sum(f(r['ZundCost']) + f(r['CrestCost']) + f(r['SaberCost']) + f(r['BobstCost']) for r in rows),
            'Strip': sum(f(r['HandStripCost']) for r in rows),
            'Glue': sum(f(r['GluingCost']) for r in rows),
            'Freight': sum(f(r['FreightCost']) for r in rows)}


def mk(est, iq, job, qty, opt, s_rows, i_rows, note=''):
    proj = s_rows[0]['Project Name'] if s_rows else ''
    sc, ic = spot_cats(s_rows), iq_cats(i_rows)
    return {'est': est, 'iq': iq, 'job': job, 'qty': qty, 'opt': opt, 'note': note, 'proj': proj,
            's_rows': s_rows, 'i_rows': i_rows,
            'spot_sell': sum(f(r['Sell Total']) for r in s_rows),
            'iq_sell': sum(f(r['ComponentTotalPrice']) for r in i_rows),
            'spot_cost': sum(sc.values()),
            'iq_cost': sum(ic.values()),
            'scats': sc, 'icats': ic}


S = defaultdict(list)
for r in spot_rows:
    S[r['SPOT Estimate #']].append(r)
I = defaultdict(list)
for r in iq_rows:
    I[(r['IdQuote'], r['OptDesc'])].append(r)


def srows(est, qty=None, comp=None):
    return [r for r in S[est] if (qty is None or r['Qty'] == qty) and (comp is None or comp in r['Component'])]


def irows(qid, opt, qty=None, prod_not=None):
    return [r for r in I[(qid, opt)] if (qty is None or r['Quantity'] == qty)
            and (prod_not is None or prod_not not in str(r['Product']))]


matches = []
matches.append(mk('RW-260706-094426', 223488, 'No', '300', 'CAD 1-Side (mailer only)',
                  srows('RW-260706-094426', comp='Digital Print outside only'),
                  irows(223488, 'CAD Cut Print 1 Side', 300), 'per-component; 4 inserts have no iQuote counterpart'))
matches.append(mk('RW-260706-094426', 223488, 'No', '300', 'CAD 2-Sides (mailer only)',
                  srows('RW-260706-094426', comp='outside AND inside'),
                  irows(223488, 'CAD Cut Print 2 Sides', 300), 'per-component'))
matches.append(mk('RW-260720-082503', 223583, 'Yes', '918+1750', 'Option 1 (kit)',
                  srows('RW-260720-082503'), irows(223583, 'Option 1'),
                  'SPOT: 1 comp x 2 qty; iQuote: 2 products 919+1751'))
matches.append(mk('RW-260629-133153', 223378, 'No', '250 sets', 'CadCut (kit)',
                  srows('RW-260629-133153'), irows(223378, 'CadCut'),
                  'SPOT side panel qty 500 = 250 sets of 2'))
matches.append(mk('RW-260709-132147', 221998, 'No', '525 kit', 'D/C w/riser minus Riser',
                  srows('RW-260709-132147'), irows(221998, 'Option 1 - D/C w/riser', prod_not='Riser'),
                  'Riser excluded (not in SPOT est)'))
MANUAL = {'RW-260706-094426', 'RW-260720-082503', 'RW-260629-133153', 'RW-260709-132147'}

iq_scen = defaultdict(list)
for r in iq_rows:
    iq_scen[(r['IdQuote'], r['OptDesc'], round(f(r['TotalEstimatePrice']), 2))].append(r)


def qm(a, b):
    return a == b or abs(a - b) <= max(2, 0.02 * max(a, b))


for est, rows in S.items():
    if est in MANUAL:
        continue
    comps = set(r['Component'] for r in rows)
    qtys = sorted(set(r['Qty'] for r in rows))
    kit = len(rows) == len(comps) and len(qtys) > 1
    groups = [(f"{'/'.join(map(str,qtys))} kit", rows)] if kit else [(str(q), [r for r in rows if r['Qty'] == q]) for q in qtys]
    routes = set((r['Route'] or '') for r in rows)
    s_route = 'CAD' if routes <= {'CAD', ''} else 'Die'
    for qlabel, grp in groups:
        gq = set(r['Qty'] for r in grp)
        cands = []
        for (qid, opt, tep), ig in iq_scen.items():
            if qid != grp[0]['iQuote #']:
                continue
            iqq = set(r['Quantity'] for r in ig)
            if all(any(qm(x, y) for y in iqq) for x in gq):
                zund = sum(f(r['ZundCost']) for r in ig)
                cut = sum(f(r['ZundCost']) + f(r['CrestCost']) + f(r['SaberCost']) + f(r['BobstCost']) for r in ig)
                i_route = 'CAD' if zund >= cut - zund else 'Die'
                cands.append((opt, ig, i_route))
        if not cands:
            continue
        pref = [c for c in cands if c[2] == s_route] or cands
        s_sell = sum(f(r['Sell Total']) for r in grp)
        opt, ig, _ = min(pref, key=lambda c: abs(sum(f(r['ComponentTotalPrice']) for r in c[1]) - s_sell))
        matches.append(mk(est, grp[0]['iQuote #'], grp[0]['Became a Job'], qlabel, opt, grp, ig))

# produced jobs: (mis_job, spot_est, spot_qty, iquote_id, option, iq_qty, note)
JOBS = [
    (236341, 'RW-260623-091031', 210, 223207, 'Die Cut', 210, 'Job ran DIE CUT; SPOT quoted CAD route'),
    (236355, 'RW-260727-080830', 1, 223209, 'Option 1', 1, ''),
    (236550, 'RW-260727-102959', 230, 223393, 'Option 1', 230, ''),
    (236688, 'RW-260727-095102', 50, 223486, 'Option 1', 50, ''),
    (236802, 'RW-260727-101129', 220, 223454, 'Option 1 - D/C', 220, ''),
    (236822, 'RW-260727-144121', 500, 223575, 'Option 1', 500, ''),
    (236943, 'RW-260722-082345', 250, 223591, 'CAD Cut - for pricing purposes', 250, 'Sold at CAD-pricing option; SPOT quoted Die route'),
    (236869, 'RW-260730-155412', 250, 223597, 'Option 1', 250, 'Invoiced $2,876.15 vs quoted $2,749.42'),
    (237252, 'RW-260724-113737', 22, 223651, 'Option 1', 22, 'Cut-only op, no material'),
]


def spot_cat(est, qty):
    rows = [r for r in spot_rows if r['SPOT Estimate #'] == est and r['Qty'] == qty]
    return {'Material': sum(f(r['Material Cost']) for r in rows),
            'Print': sum(f(r['Print Cost']) for r in rows),
            'Cutting': sum(f(r['Cutting Cost']) for r in rows),
            'Hand': sum(f(r['Strip Cost']) + f(r['Glue Cost']) + f(r['Tape Cost']) for r in rows),
            'Die': sum(f(r['Die']) for r in rows),
            'Total': sum(f(r['Total Cost']) for r in rows),
            'Sell': sum(f(r['Sell Total']) for r in rows)}


def iq_cat(qid, opt, qty):
    rows = [r for r in iq_rows if r['IdQuote'] == qid and r['OptDesc'] == opt and r['Quantity'] == qty]
    return {'Material': sum(f(r['MaterialCost']) for r in rows),
            'Print': sum(f(r['PressCost']) for r in rows),
            'Cutting': sum(f(r['ZundCost']) + f(r['CrestCost']) + f(r['SaberCost']) + f(r['BobstCost']) for r in rows),
            'Hand': sum(f(r['HandStripCost']) + f(r['GluingCost']) for r in rows),
            'Die': 0.0,
            'Total': sum(f(r['TotalActualCost']) - f(r['NotInSpotCost']) for r in rows),
            'Sell': sum(f(r['ComponentTotalPrice']) for r in rows)}


def act_cat(job_id):
    r = next(a for a in act_rows if a['job_id'] == job_id)
    cut = sum(f(r[k]) for k in ['LGZundCost', 'SmZundCost', 'KongsbergCost', 'BobstCost', 'CrestClamshellCost', 'SaberCost'])
    return {'Material': f(r['MaterialCost']) + f(r['OtherMaterialCost']),
            'Print': f(r['PrintCost']),
            'Cutting': cut,
            'Hand': f(r['HandStripCost']) + f(r['HandGlueCost']) + f(r['HandTapeCost']),
            'Die': f(r['DieCost']),
            'Total': f(r['TotalActualCost']),
            'Sell': f(r['SellTotal'])}
