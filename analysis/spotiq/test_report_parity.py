import json, os, sys, io
import requests

API = None
for line in open('/app/frontend/.env'):
    if line.startswith('REACT_APP_BACKEND_URL='):
        API = line.split('=', 1)[1].strip()
TOKEN = requests.post(f"{API}/api/auth/login", json={
    "email": "admin@standardprinting.com", "password": "spot-admin-2026"}).json()["access_token"]
H = {"Authorization": f"Bearer {TOKEN}"}


def comp(desc, w, l, rule, strip="No"):
    return {"id": f"it-{desc}", "desc": desc, "type": "Manufactured", "pcsPerSet": 1,
            "state": {"itemDesc": desc, "oneUpW": w, "oneUpL": l,
                      "flute": "B", "boardGrade": "200", "outsideLiner": "Kemi", "insideLiner": "Kemi",
                      "printReq": "Yes", "printMode": "Gloss", "printSides": "1",
                      "cadDifficulty": "Standard", "cadPricingMethod": "Rule",
                      "totalRuleInches": rule, "handStripping": strip,
                      "strippingComplexity": "Standard",
                      "prepressTreatment": "Amortize", "prepressCost": "50",
                      "freightCost": "0", "freightTreatment": "None"}}


state = {
    "quoteType": "Multi-Part", "pricingMode": "Matched Set",
    "qtys": ["8", "15", "", "", ""], "activeItemIdx": 0,
    "custName": "REPORT PARITY TEST", "custRef": "Report Parity Test",
    "goodForTesting": True, "iquoteNum": "TEST-1", "iquoteWon": False,
    "bomItems": [comp("Back Panel", "55", "51.3", "637", "Yes"),
                 comp("Side Panel", "40", "30", "300"),
                 comp("Tray", "20", "16", "150")],
}
# item-level top-level mirror keys for active item 0 (as UI save does)
state.update(state["bomItems"][0]["state"])

q = requests.post(f"{API}/api/quotes", headers=H, json={"state": state, "est_num": "PARITY-TEST-1"}).json()
qid = q["id"]
print("quote:", qid)

# UI-equivalent engine call (buildMpPayload shape)
payload = {
    "quoteType": "Multi-Part", "pricingMode": "Matched Set", "setMarginPct": "0.40",
    "setMarginOverride": False, "fulfillmentMarginPct": "0.45", "purchasedMarginPct": "0.40",
    "qtys": ["8", "15", "", "", ""], "activeItemIdx": 0,
    "bomItems": state["bomItems"], "custName": state["custName"], "custRef": state["custRef"],
    "jobFreightBySlot": {}, "quoteId": qid,
}
eng = requests.post(f"{API}/api/engine/calculate-multi-part", headers=H, json=payload).json()
ui = {}
for slot in eng.get("slots", []):
    for pc in slot.get("perComponent", []):
        ui[(pc["desc"], pc["qty"])] = (round(pc["itemSell"], 2), round(slot["sellTotal"], 2),
                                       round(pc["totalMfg"], 2), round(slot["gmPct"] * 100, 1))
print("UI slots:", json.dumps({f"{k[0]}@{k[1]}": v for k, v in ui.items()}, indent=0))

# report
import openpyxl
r = requests.get(f"{API}/api/quotes/_report/testing.xlsx", headers=H)
wb = openpyxl.load_workbook(io.BytesIO(r.content))
ws = wb.active
hdr = [c.value for c in ws[1]]
ok = True
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(hdr, row))
    if d["SPOT Estimate #"] != "PARITY-TEST-1":
        continue
    key = (d["Component"], d["Qty"])
    exp = ui.get(key)
    got = (d["Sell Total"], d["Quote Total"], d["Total Cost"], d["Margin %"])
    match = exp == got
    ok = ok and match and exp is not None
    print(f"{key}: report sell/quote/cost/gm={got} ui={exp} {'MATCH' if match else 'MISMATCH!'}")
print("PARITY:", "PASS" if ok else "FAIL")
requests.delete(f"{API}/api/quotes/{qid}", headers=H)
