"""What-if on the matched print set: new Inca speeds (by finish or single) + setup minutes/pass.
Uses SPOT's own sheet counts (its layout) so the result is what customers would actually see."""
import json
import sys
from collections import defaultdict

import openpyxl

sys.path.insert(0, "/app/analysis/print")
from build_print_parity import QTY_BANDS, SP_SPEED, spot_formula, write_sheet, r2  # noqa: E402

HERE = "/app/analysis/print"
rows = [r for r in json.load(open(f"{HERE}/print_parity_data.json"))["matched"]
        if not r["iQ Ganged"] and r["Same Finish"] == "Yes" and r["iQ Print $"] > 0
        and r["SPOT Print Mode"] in ("Satin", "Gloss")]


def spot_print(r, speeds, setup_min, rate=485.0):
    mode = r["SPOT Print Mode"]
    sq = float(r["SPOT Sheet W"]) * float(r["SPOT Sheet L"]) / 144
    return spot_formula(r["SPOT Sheets"], sq, mode, r["SPOT White Ink"] == "Yes", r["SPOT Sides"] == "2-Sided",
                        speed=speeds[mode], setup_min=setup_min, rate=rate)["total"]


SCENARIOS = {
    "A. Today 75-75 @485, no setup": ({"Satin": 75, "Gloss": 75}, 0, 485),
    "B. Single 160 @485, no setup": ({"Satin": 160, "Gloss": 160}, 0, 485),
    "C. By finish 190-135 @485, no setup": ({"Satin": 190, "Gloss": 135}, 0, 485),
    "D. By finish 190-135 @485 +10min": ({"Satin": 190, "Gloss": 135}, 10, 485),
    "E. By finish 220-160 @485 +5min": ({"Satin": 220, "Gloss": 160}, 5, 485),
    "F. iQuote 144-102 @319 +10min": ({"Satin": 144, "Gloss": 102}, 10, 319),
    "G. iQuote 144-102 @319 +7min": ({"Satin": 144, "Gloss": 102}, 7, 319),
    "H. iQuote 144-102 @319 +5min": ({"Satin": 144, "Gloss": 102}, 5, 319),
}


def band_table(fn):
    g = defaultdict(list)
    for r in rows:
        g[r["Qty Band"]].append(r)
    out = []
    for lab, *_ in QTY_BANDS:
        rs = g.get(lab)
        if not rs:
            continue
        sp = sum(fn(r) for r in rs)
        iq = sum(r["iQ Print $"] for r in rs)
        pcts = sorted((fn(r) - r["iQ Print $"]) / r["iQ Print $"] * 100 for r in rs)
        mape = sum(abs(p) for p in pcts) / len(pcts)
        out.append({"Qty Band": lab, "Rows": len(rs), "SPOT $": r2(sp), "iQ $": r2(iq),
                    "Δ % agg": round((sp - iq) / iq * 100, 1), "Δ % median": round(pcts[len(pcts) // 2], 1),
                    "MAPE %": round(mape, 1), "Rows within ±20%": sum(1 for p in pcts if abs(p) <= 20)})
    return out


def finish_table(fn):
    out = []
    for fin in ("Satin", "Gloss"):
        rs = [r for r in rows if r["SPOT Print Mode"] == fin]
        sp = sum(fn(r) for r in rs); iq = sum(r["iQ Print $"] for r in rs)
        pcts = [(fn(r) - r["iQ Print $"]) / r["iQ Print $"] * 100 for r in rs]
        out.append({"Finish": fin, "Rows": len(rs), "SPOT $": r2(sp), "iQ $": r2(iq),
                    "Δ % agg": round((sp - iq) / iq * 100, 1), "MAPE %": round(sum(abs(p) for p in pcts) / len(pcts), 1)})
    return out


def main():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    overview = []
    for name, (speeds, setup, rate) in SCENARIOS.items():
        fn = lambda r, s=speeds, m=setup, rt=rate: spot_print(r, s, m, rt)  # noqa: E731
        bt = band_table(fn); ft = finish_table(fn)
        sp = sum(fn(r) for r in rows); iq = sum(r["iQ Print $"] for r in rows)
        pcts = [(fn(r) - r["iQ Print $"]) / r["iQ Print $"] * 100 for r in rows]
        overview.append({"Scenario": name, "Satin sh/hr": speeds["Satin"], "Gloss sh/hr": speeds["Gloss"], "Setup min/pass": setup, "$/hr": rate,
                         "Δ % all rows": round((sp - iq) / iq * 100, 1), "MAPE % all": round(sum(abs(p) for p in pcts) / len(pcts), 1),
                         **{f"Δ% {b['Qty Band']}": b["Δ % agg"] for b in bt},
                         **{f"Δ% {f['Finish']}": f["Δ % agg"] for f in ft},
                         "Rows within ±20%": sum(1 for p in pcts if abs(p) <= 20), "Rows": len(rows)})
        print(f"\n{name}"); [print("  ", b) for b in bt]; [print("  ", f) for f in ft]
        write_sheet(wb, name[:31].replace("/", "-"), bt + [{}] + ft)
    ws = wb.create_sheet("Overview", 0)
    write_sheet(wb, "_tmp", overview)
    tmp = wb["_tmp"]
    for row in tmp.iter_rows(values_only=True):
        ws.append(list(row))
    wb.remove(tmp)
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    ws.column_dimensions["A"].width = 42
    wb.save(f"{HERE}/Print_Speed_Scenarios.xlsx")
    print("\nOVERVIEW"); [print(o) for o in overview]


if __name__ == "__main__":
    main()
