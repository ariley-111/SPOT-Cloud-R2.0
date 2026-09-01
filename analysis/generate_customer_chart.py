"""Chart-only workbook: Top 15 Customers × Qty Band (Median Variance %).

Two views: heatmap-style color-scaled table + a grouped horizontal bar chart.

Usage: python3 generate_customer_chart.py <round2_input.xlsx> <output.xlsx>
"""
import sys
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule


IN  = sys.argv[1] if len(sys.argv) > 1 else "input.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "SPOT_Customer_QtyBand_Chart.xlsx"

# Rebuild the quote-total dataset (same rule as Round 2 v2)
iq = pd.read_excel(IN, sheet_name="iQuote Data")
sp = pd.read_excel(IN, sheet_name="SPOT Data")

def norm_opt(x):
    if pd.isna(x): return ""
    s = str(x).strip().lower()
    if " - " in s: s = s.split(" - ")[0].strip()
    return s

iq["_opt"] = iq["OptDesc"].apply(norm_opt)
sp["_opt"] = sp["iQuote Option"].apply(norm_opt)

groups = sp.groupby(["iQuote Estimate #","Quote Total"], dropna=False)
rows = []
for (est, qt), g in groups:
    pairs = set(zip(g["_opt"], g["Qty"]))
    iq_pool = iq[iq["IdQuote"] == est]
    matched = iq_pool[iq_pool.apply(lambda r: (r["_opt"], r["Quantity"]) in pairs, axis=1)]
    if len(matched) == 0:
        matched = iq_pool[iq_pool["Quantity"].isin(set(g["Qty"].tolist()))]
    if len(matched) == 0:
        matched = iq_pool
    iq_sum = matched["ComponentTotalPrice"].sum() if len(matched) > 0 else np.nan
    rows.append({
        "Customer": g["Customer"].iloc[0],
        "quote_total": qt,
        "iq_sum": iq_sum,
        "qty_max": g["Qty"].max(),
    })
df = pd.DataFrame(rows)
df["variance_pct"] = (df["quote_total"] - df["iq_sum"]) / df["iq_sum"] * 100

def qty_band(q):
    if pd.isna(q): return "Unknown"
    q = int(q)
    if q < 50: return "1-49"
    if q < 200: return "50-199"
    if q < 1000: return "200-999"
    if q < 5000: return "1000-4999"
    return "5000+"

df["Qty Band"] = df["qty_max"].apply(qty_band)

BANDS = ["1-49","50-199","200-999","1000-4999","5000+"]
top15 = df.groupby("Customer").size().sort_values(ascending=False).head(15).index.tolist()

sub = df[df["Customer"].isin(top15)]
pivot = sub.pivot_table(index="Customer", columns="Qty Band", values="variance_pct", aggfunc="median")
pivot = pivot.reindex(columns=BANDS).reindex(top15)
counts = sub.groupby("Customer").size().reindex(top15)


# ─── Build workbook ─────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Customer × Qty Band"

TITLE = Font(bold=True, size=16, color="1F4E78")
SUB   = Font(italic=True, size=10, color="595959")
HDR   = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
BOLD  = Font(bold=True)
THIN  = Side(border_style="thin", color="BFBFBF")
BORDER= Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER= Alignment(horizontal="center", vertical="center")

FMT_PCT = "0.00\"%\";-0.00\"%\";0.00\"%\""
FMT_INT = "#,##0"

ws["A1"] = "Median Variance % — Top 15 Customers × Qty Band"
ws["A1"].font = TITLE
ws["A2"] = ("Quote-total level. Cells show median variance (SPOT vs iQuote). "
            "Green = SPOT cheaper, Red = SPOT more expensive, White = near parity. "
            "Blank cell = no quotes from that customer in that qty band.")
ws["A2"].font = SUB
ws.merge_cells("A1:H1")
ws.merge_cells("A2:H2")

# Header row (row 4)
headers = ["Customer"] + BANDS + ["Quote Count"]
for j, h in enumerate(headers, start=1):
    c = ws.cell(row=4, column=j, value=h)
    c.font = HDR; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER

# Data rows (5 onward)
for i, cust in enumerate(top15, start=5):
    ws.cell(row=i, column=1, value=cust).font = BOLD
    ws.cell(row=i, column=1).border = BORDER
    for j, band in enumerate(BANDS, start=2):
        v = pivot.loc[cust, band]
        c = ws.cell(row=i, column=j, value=None if pd.isna(v) else float(v))
        c.number_format = FMT_PCT
        c.border = BORDER
        c.alignment = CENTER
    n_cell = ws.cell(row=i, column=7, value=int(counts.loc[cust]))
    n_cell.number_format = FMT_INT
    n_cell.border = BORDER
    n_cell.alignment = CENTER

# Heatmap conditional formatting on the band cells (B5:F19)
rule = ColorScaleRule(
    start_type="num", start_value=-30, start_color="63BE7B",  # green
    mid_type="num",   mid_value=0,     mid_color="FFFFFF",     # white
    end_type="num",   end_value=60,    end_color="F8696B",     # red
)
ws.conditional_formatting.add(f"B5:F{4+len(top15)}", rule)

# Column widths
ws.column_dimensions["A"].width = 42
for c in "BCDEFG":
    ws.column_dimensions[c].width = 12
ws.freeze_panes = "B5"

# ─── Grouped bar chart (horizontal) ─────────────────────────────
chart = BarChart()
chart.type = "bar"
chart.style = 11
chart.grouping = "clustered"
chart.title = "Median Variance % — by Customer × Qty Band"
chart.x_axis.title = "Median Variance %"
chart.y_axis.title = "Customer"
chart.height = 22
chart.width = 32
chart.legend.position = "b"

# Data: columns B (1-49) through F (5000+), rows 4 (header) → 4+n
data_ref = Reference(ws, min_col=2, max_col=6, min_row=4, max_row=4 + len(top15))
cats_ref = Reference(ws, min_col=1, min_row=5, max_row=4 + len(top15))
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
ws.add_chart(chart, "I4")

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Customers: {len(top15)}, Bands: {len(BANDS)}")
